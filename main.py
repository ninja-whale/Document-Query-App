import streamlit as st
import sqlite3
from passlib.hash import bcrypt
import PyPDF2
from docx import Document
import io
import pandas as pd
from fpdf import FPDF
from get_query import get_query_main


# Database connection
conn = sqlite3.connect('users.db')
c = conn.cursor()

# Create necessary tables
c.execute('''CREATE TABLE IF NOT EXISTS users
             (username TEXT PRIMARY KEY, password TEXT)''')

c.execute('''CREATE TABLE IF NOT EXISTS user_history
             (username TEXT, query TEXT, response TEXT, timestamp DATETIME DEFAULT CURRENT_TIMESTAMP)''')

# Function for the login and registration page
def auth_page():
    option = st.sidebar.selectbox('Choose an option', ['Login', 'Register'])
    if option == 'Register':
        st.subheader('Register')
        username = st.text_input('Username', key='register_username')
        password = st.text_input('Password', type='password', key='register_password')
        if st.button('Register'):
            if username and password:
                create_user(username, password)
                st.success('Registration successful!')

    elif option == 'Login':
        st.subheader('Login')
        username = st.text_input('Username', key='login_username')
        password = st.text_input('Password', type='password', key='login_password')
        if st.button('Login'):
            if username and password:
                if verify_user(username, password):
                    st.success(f'Welcome, {username}!, please click on login to continue.')
                    st.session_state['authenticated'] = True
                    st.session_state['username'] = username
                else:
                    st.error('Invalid credentials, you may register first')

# Function for the document query page
def query_page():
    if 'authenticated' in st.session_state and st.session_state['authenticated']:
        st.subheader('Upload a file')
        uploaded_file = st.file_uploader("Choose a file", type=["pdf", "txt", "docx"])
        if uploaded_file is not None:
            if uploaded_file.type == "application/pdf":
                extracted_text = extract_text_from_pdf(uploaded_file)
            elif uploaded_file.type == "text/plain":
                extracted_text = extract_text_from_txt(uploaded_file)
            elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                extracted_text = extract_text_from_docx(uploaded_file)
            else:
                st.error("Unsupported file format")
                return

            st.text_area("Extracted Text", extracted_text, height=200)
            query = st.text_input("Ask a query about the text:")
            if st.button("Search"):
                result = get_query_main(query, extracted_text)
                log_user_query(st.session_state['username'], query, result)
                st.success(f'{result}')

        st.subheader('Download Your Chat History')
        download_chat_history_pdf(st.session_state['username'])

        st.subheader('View Your Query History')
        show_user_history(st.session_state['username'])
    else:
        st.warning("Please log in to access this page.")

# Function for the logout page
def logout_page():
    st.title("You have been logged out.")
    st.write("Thank you for using the app.")
    # Clear session state
    for key in list(st.session_state.keys()):
        del st.session_state[key]
    st.sidebar.success("Logged out successfully. Please refresh and log in again to continue.")

# Function to log user query
def log_user_query(username, query, response):
    c.execute('INSERT INTO user_history (username, query, response) VALUES (?, ?, ?)', (username, query, response))
    conn.commit()

# Function to download chat history as PDF
def download_chat_history_pdf(username):
    query = "SELECT query, response, timestamp FROM user_history WHERE username = ?"
    df = pd.read_sql(query, conn, params=(username,))
    if df.empty:
        st.warning("No chat history available.")
        return
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for index, row in df.iterrows():
        pdf.cell(0, 10, f"Query: {row['query']}", ln=True)
        pdf.cell(0, 10, f"Response: {row['response']}", ln=True)
        pdf.cell(0, 10, f"Timestamp: {row['timestamp']}", ln=True)
        pdf.ln(10)
    buf = io.BytesIO()
    pdf_output = pdf.output(dest='S').encode('latin1')
    buf.write(pdf_output)
    buf.seek(0)
    st.download_button(label="Download Chat History as PDF", data=buf, file_name=f"{username}_chat_history.pdf", mime="application/pdf")

# Function to show user history
def show_user_history(username):
    query = "SELECT query, response, timestamp FROM user_history WHERE username = ?"
    df = pd.read_sql(query, conn, params=(username,))
    if df.empty:
        st.warning("No chat history available.")
    else:
        st.subheader("Your Query History")
        st.write(df)

# Authentication Functions
def create_user(username, password):
    hashed_password = bcrypt.hash(password)
    c.execute('INSERT INTO users (username, password) VALUES (?, ?)', (username, hashed_password))
    conn.commit()

def verify_user(username, password):
    c.execute('SELECT password FROM users WHERE username = ?', (username,))
    stored_password = c.fetchone()
    if stored_password and bcrypt.verify(password, stored_password[0]):
        return True
    return False

# Text Extraction Functions
def extract_text_from_pdf(file):
    pdf_reader = PyPDF2.PdfReader(file)
    text = ""
    for page in range(len(pdf_reader.pages)):
        text += pdf_reader.pages[page].extract_text()
    return text

def extract_text_from_txt(file):
    return file.read().decode("utf-8")

def extract_text_from_docx(file):
    doc = Document(file)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

# Main function to handle navigation
def main():
    if 'authenticated' not in st.session_state or not st.session_state['authenticated']:
        auth_page()  # Show the login/register page first if not authenticated
    else:
        st.sidebar.title("Navigation")
        page = st.sidebar.radio("Go to", ["Query Document", "Logout"])
        if page == "Query Document":
            query_page()
        elif page == "Logout":
            logout_page()

if __name__ == '__main__':
    main()
